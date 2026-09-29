"""Parse Road Formation Level workbook into frontend/public/road_formation.json.

Layout (Sheet1): three side-by-side blocks per row —
  LHS 30M   (cols B..G): Latitude, Longitude, Chainage, Elevation, UTM Zone, Min Formation Level
  Centerline(cols I..M): Latitude, Longitude, Chainage, Elevation, Min Formation Level
  RHS 30M   (cols O..S): Latitude, Longitude, Chainage, Elevation, Min Formation Level

Note: the workbook's LHS / Centerline lat-lon columns are scrambled along the
corridor (path length ~7× chainage, regular teleports every 64 rows). RHS is
spatially continuous and matches the ~34.5 km corridor, so map geometry for
LHS and Centerline is rebuilt as 60 m / 30 m left-offsets of the RHS spine.
Formation levels and elevations are kept from the original columns.
"""
from __future__ import annotations

import json
import math
import sys
from pathlib import Path

import openpyxl

SRC = Path(r"C:\Users\Kunal.Desale\Downloads\Road Formation Level 2.xlsx")
OUT = Path(r"C:\Users\Kunal.Desale\Desktop\Adani\frontend\public\road_formation.json")

# (branch id, name, column offsets: lat, lon, chainage, elevation, formation)
BLOCKS = [
    ("lhs", "LHS 30M", 1, 2, 3, 4, 6),
    ("centerline", "Centerline", 8, 9, 10, 11, 12),
    ("rhs", "RHS 30M", 14, 15, 16, 17, 18),
]

# Distance of each branch left of the RHS spine (RHS is +30 m from centerline).
OFFSET_M = {
    "rhs": 0.0,
    "centerline": -30.0,
    "lhs": -60.0,
}


def _num(v) -> float | None:
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _bearing_rad(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    φ1, φ2 = math.radians(lat1), math.radians(lat2)
    Δλ = math.radians(lon2 - lon1)
    y = math.sin(Δλ) * math.cos(φ2)
    x = math.cos(φ1) * math.sin(φ2) - math.sin(φ1) * math.cos(φ2) * math.cos(Δλ)
    return math.atan2(y, x)


def _offset(lat: float, lon: float, bearing_rad: float, distance_m: float) -> tuple[float, float]:
    """Move `distance_m` along `bearing_rad` from (lat, lon). Negative = opposite."""
    if abs(distance_m) < 1e-9:
        return lat, lon
    R = 6_371_000.0
    δ = distance_m / R
    θ = bearing_rad
    φ1 = math.radians(lat)
    λ1 = math.radians(lon)
    φ2 = math.asin(
        math.sin(φ1) * math.cos(δ) + math.cos(φ1) * math.sin(δ) * math.cos(θ)
    )
    λ2 = λ1 + math.atan2(
        math.sin(θ) * math.sin(δ) * math.cos(φ1),
        math.cos(δ) - math.sin(φ1) * math.sin(φ2),
    )
    return math.degrees(φ2), math.degrees(λ2)


def _path_length_km(points: list[dict]) -> float:
    total = 0.0
    for i in range(1, len(points)):
        a, b = points[i - 1], points[i]
        dlat = (a["lat"] - b["lat"]) * 111_000
        dlon = (a["lon"] - b["lon"]) * 111_000 * math.cos(math.radians(a["lat"]))
        total += math.hypot(dlat, dlon)
    return total / 1000.0


def _rebuild_offsets_from_rhs(branches: dict[str, dict]) -> None:
    """Replace LHS / Centerline coordinates with offsets of the continuous RHS spine."""
    rhs = branches["rhs"]["points"]
    if len(rhs) < 2:
        return

    bearings: list[float] = []
    for i, p in enumerate(rhs):
        if i < len(rhs) - 1:
            bearings.append(_bearing_rad(p["lat"], p["lon"], rhs[i + 1]["lat"], rhs[i + 1]["lon"]))
        else:
            bearings.append(bearings[-1])

    for bid, distance_m in OFFSET_M.items():
        if bid == "rhs" or bid not in branches:
            continue
        pts = branches[bid]["points"]
        n = min(len(pts), len(rhs))
        for i in range(n):
            # Left of travel direction = bearing − 90°.
            left = bearings[i] - math.pi / 2
            lat, lon = _offset(rhs[i]["lat"], rhs[i]["lon"], left, abs(distance_m))
            pts[i]["lat"] = lat
            pts[i]["lon"] = lon


def main() -> None:
    if not SRC.exists():
        print(f"Missing workbook: {SRC}", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(SRC, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]

    branches: dict[str, dict] = {
        bid: {"id": bid, "name": name, "points": []} for bid, name, *_ in BLOCKS
    }

    for row in ws.iter_rows(min_row=4, values_only=True):
        for bid, _name, c_lat, c_lon, c_ch, c_el, c_fl in BLOCKS:
            lat = _num(row[c_lat]) if len(row) > c_lat else None
            lon = _num(row[c_lon]) if len(row) > c_lon else None
            ch = _num(row[c_ch]) if len(row) > c_ch else None
            if lat is None or lon is None or ch is None:
                continue
            branches[bid]["points"].append(
                {
                    "chainage_km": ch,
                    "lat": lat,
                    "lon": lon,
                    "ground_elev_m": _num(row[c_el]) if len(row) > c_el else None,
                    "formation_level_m": _num(row[c_fl]) if len(row) > c_fl else None,
                }
            )

    # Align all branches to the same row count (RHS is the geographic authority).
    n = min(len(branches[bid]["points"]) for bid, *_ in BLOCKS)
    for bid, *_ in BLOCKS:
        branches[bid]["points"] = branches[bid]["points"][:n]

    before = {bid: round(_path_length_km(branches[bid]["points"]), 2) for bid, *_ in BLOCKS}
    _rebuild_offsets_from_rhs(branches)
    after = {bid: round(_path_length_km(branches[bid]["points"]), 2) for bid, *_ in BLOCKS}
    print("path length km before:", before)
    print("path length km after: ", after)

    for b in branches.values():
        b["count"] = len(b["points"])

    all_pts = [p for b in branches.values() for p in b["points"]]
    fls = [p["formation_level_m"] for p in all_pts if p["formation_level_m"] is not None]
    els = [p["ground_elev_m"] for p in all_pts if p["ground_elev_m"] is not None]
    chs = [p["chainage_km"] for p in all_pts]

    payload = {
        "title": "Computed Road Formation Level (per Schedule-B, Annexure-I, Cl.4)",
        "source_file": SRC.name,
        "count": len(all_pts),
        "from_km": min(chs) if chs else None,
        "to_km": max(chs) if chs else None,
        "formation_min_m": min(fls) if fls else None,
        "formation_max_m": max(fls) if fls else None,
        "ground_min_m": min(els) if els else None,
        "ground_max_m": max(els) if els else None,
        "geometry_note": (
            "LHS and Centerline map coordinates are rebuilt as 60 m / 30 m "
            "left-offsets of the RHS spine; formation levels come from the workbook."
        ),
        "branches": [branches[bid] for bid, *_ in BLOCKS],
    }

    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, separators=(",", ":"))

    for b in payload["branches"]:
        print(f"{b['name']}: {b['count']} points")
    print(
        f"chainage {payload['from_km']}–{payload['to_km']} km · "
        f"formation {payload['formation_min_m']}–{payload['formation_max_m']} m · "
        f"ground {payload['ground_min_m']}–{payload['ground_max_m']} m"
    )
    print("written", OUT)


if __name__ == "__main__":
    main()
