"""Parse the excavation earthwork workbook → frontend/public/cut_fill.json.

Schema (Excavation Datasheet New 2.xlsx):
  Sheets: Summary (title only), RHS, LHS
  Columns per side row:
    Latitude, Longitude, Chainage From (km), Chainage To (km), Length (m),
    Center OGL, OGL at 30m <side>, Avg OGL, Finished Road Level,
    Avg Fill Height (m), 30m Fill Vol (m3), 30m Fill Mass (T)

This dataset is fill-only (no cut, no TCS). The RHS latitude/longitude track
chainage cleanly, so they are used as the positional anchor for BOTH sides
(the LHS coordinates in the sheet are scrambled). Each aggregated segment
carries its own lat/lon so the map can place points directly.

Cut volumes are merged in from the earlier earthwork workbook
(01a4eab9... .xlsx, stashed as data/cut_fill_cut.xlsx), which carries
"Cut Volume (m3)" / "Cut Mass (T)" per 10 m station on both sides.
"""
from __future__ import annotations

import json
import os
import shutil
from collections import Counter

import openpyxl

ROOT = os.path.dirname(__file__)
DEFAULT_SRC = os.path.join(ROOT, "data", "cut_fill.xlsx")
DOWNLOADS_CANDIDATES = [
    r"C:\Users\Kunal.Desale\Downloads\Excavation Datasheet New 2.xlsx",
]
# Older earthwork workbook that still carries per-station cut volumes.
CUT_SRC = os.path.join(ROOT, "data", "cut_fill_cut.xlsx")
CUT_DOWNLOADS_CANDIDATES = [
    r"C:\Users\Kunal.Desale\Downloads\01a4eab92e1542f9b160e14bfcd068f9 (1).xlsx",
    r"C:\Users\Kunal.Desale\Downloads\01a4eab92e1542f9b160e14bfcd068f9.xlsx",
]
OUT = os.path.join(ROOT, "..", "frontend", "public", "cut_fill.json")

# Aggregate 10 m stations into map points of this length (km).
AGG_KM = 0.1
DENSITY_T_PER_M3 = 1.8


def _num(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _classify(height_m: float | None) -> str:
    """Fill-height class for corridor colouring (fill-only dataset)."""
    h = height_m if height_m is not None else 0.0
    if h < 2:
        return "fill_0_2"
    if h < 4:
        return "fill_2_4"
    if h < 6:
        return "fill_4_6"
    if h < 8:
        return "fill_6_8"
    return "fill_8_plus"


def _read_side(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    col = {h: i for i, h in enumerate(header)}

    def g(row, *names):
        for n in names:
            if n in col:
                return row[col[n]]
        return None

    out: list[dict] = []
    for row in rows[1:]:
        if not row:
            continue
        frm = _num(g(row, "Chainage From (km)"))
        to = _num(g(row, "Chainage To (km)"))
        if frm is None or to is None:
            continue  # skips the TOTAL footer row
        out.append(
            {
                "lat": _num(g(row, "Latitude")),
                "lon": _num(g(row, "Longitude")),
                "from_km": round(frm, 4),
                "to_km": round(to, 4),
                "length_m": _num(g(row, "Length (m)")) or round((to - frm) * 1000, 2),
                "center_ogl_m": _num(g(row, "Center OGL")),
                "ogl_30m_m": _num(g(row, "OGL at 30m RHS", "OGL at 30m LHS")),
                "avg_ogl_m": _num(g(row, "Avg OGL")),
                "rfl_m": _num(g(row, "Finished Road Level")),
                "fill_height_m": _num(g(row, "Avg Fill Height (m)")) or 0.0,
                "fill_m3": _num(g(row, "30m Fill Vol (m³)", "30m Fill Vol (m3)")) or 0.0,
                "fill_mass_t": _num(g(row, "30m Fill Mass (T)")) or 0.0,
                "cut_m3": 0.0,
                "cut_mass_t": 0.0,
            }
        )
    return out


def _read_cut_lookup(path: str) -> dict[str, dict[float, tuple[float, float]]]:
    """side → {chainage from_km → (cut_m3, cut_mass_t)} from the old workbook."""
    lut: dict[str, dict[float, tuple[float, float]]] = {"rhs": {}, "lhs": {}}
    if not os.path.isfile(path):
        return lut
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    for sheet, side in (("RHS", "rhs"), ("LHS", "lhs")):
        if sheet not in wb.sheetnames:
            continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(h).strip() if h is not None else "" for h in rows[0]]
        col = {h: i for i, h in enumerate(header)}
        frm_i = col.get("Chainage From (km)")
        cut_i = col.get("Cut Volume (m3)", col.get("Cut Volume (m³)"))
        mass_i = col.get("Cut Mass (T)")
        if frm_i is None or cut_i is None:
            continue
        for row in rows[1:]:
            frm = _num(row[frm_i])
            cut = _num(row[cut_i])
            if frm is None or not cut:
                continue
            mass = _num(row[mass_i]) if mass_i is not None else None
            lut[side][round(frm, 3)] = (cut, mass if mass else cut * DENSITY_T_PER_M3)
    wb.close()
    return lut


def _apply_cut(rows: list[dict], cut_lut: dict[float, tuple[float, float]]) -> None:
    for r in rows:
        hit = cut_lut.get(round(r["from_km"], 3))
        if hit:
            r["cut_m3"], r["cut_mass_t"] = hit


def _coord_lookup(rows: list[dict]) -> dict[float, tuple[float, float]]:
    """chainage from_km → (lat, lon), using rows with valid coordinates."""
    lut: dict[float, tuple[float, float]] = {}
    for r in rows:
        if r["lat"] is not None and r["lon"] is not None:
            lut[round(r["from_km"], 3)] = (r["lat"], r["lon"])
    return lut


def _aggregate(
    rows: list[dict],
    coords: dict[float, tuple[float, float]],
    id_prefix: str,
    step_km: float = AGG_KM,
) -> list[dict]:
    if not rows:
        return []
    bins: dict[int, list[dict]] = {}
    for r in rows:
        idx = int(r["from_km"] / step_km + 1e-9)
        bins.setdefault(idx, []).append(r)

    segs = []
    for idx in sorted(bins):
        group = sorted(bins[idx], key=lambda r: r["from_km"])
        fill = sum(r["fill_m3"] for r in group)
        cut = sum(r["cut_m3"] for r in group)
        fill_mass = sum(r["fill_mass_t"] for r in group)
        heights = [r["fill_height_m"] for r in group if r["fill_height_m"] is not None]
        h = sum(heights) / len(heights) if heights else None
        from_km = group[0]["from_km"]
        to_km = group[-1]["to_km"]
        mid_km = round((from_km + to_km) / 2, 4)

        # Position: nearest station chainage that has clean coordinates.
        pos = None
        best = None
        for r in group:
            ll = coords.get(round(r["from_km"], 3))
            if ll is None:
                continue
            d = abs(r["from_km"] - mid_km)
            if best is None or d < best:
                best = d
                pos = ll
        if pos is None:
            # fall back to closest chainage in the whole lookup
            if coords:
                key = min(coords, key=lambda k: abs(k - mid_km))
                pos = coords[key]

        segs.append(
            {
                "id": f"{id_prefix}{idx + 1:04d}",
                "from_km": from_km,
                "to_km": to_km,
                "mid_km": mid_km,
                "lat": round(pos[0], 7) if pos else None,
                "lon": round(pos[1], 7) if pos else None,
                "length_m": round(sum(r["length_m"] for r in group), 2),
                "fill_m3": round(fill, 1),
                "cut_m3": round(cut, 1),
                "net_m3": round(fill - cut, 1),
                "fill_mass_t": round(fill_mass, 1),
                "height_cl_m": round(h, 3) if h is not None else None,
                "station_count": len(group),
                "class_id": _classify(h),
            }
        )
    return segs


def _merge_class_stretches(segs: list[dict], id_prefix: str = "CFS") -> list[dict]:
    if not segs:
        return []
    stretches: list[dict] = []
    cur = dict(segs[0])
    cur["segment_ids"] = [segs[0]["id"]]
    for s in segs[1:]:
        gap = s["from_km"] - cur["to_km"]
        same = s["class_id"] == cur["class_id"] and gap <= AGG_KM * 1.5
        if same:
            cur["to_km"] = s["to_km"]
            cur["fill_m3"] = round(cur["fill_m3"] + s["fill_m3"], 1)
            cur["cut_m3"] = round(cur["cut_m3"] + s["cut_m3"], 1)
            cur["net_m3"] = round(cur["fill_m3"] - cur["cut_m3"], 1)
            cur["fill_mass_t"] = round(cur["fill_mass_t"] + s["fill_mass_t"], 1)
            cur["length_m"] = round(cur["length_m"] + s["length_m"], 2)
            cur["station_count"] += s["station_count"]
            cur["segment_ids"].append(s["id"])
            heights = [
                x
                for x in (cur.get("height_cl_m"), s.get("height_cl_m"))
                if x is not None
            ]
            if heights:
                cur["height_cl_m"] = round(sum(heights) / len(heights), 3)
        else:
            cur["mid_km"] = round((cur["from_km"] + cur["to_km"]) / 2, 4)
            stretches.append(cur)
            cur = dict(s)
            cur["segment_ids"] = [s["id"]]
    cur["mid_km"] = round((cur["from_km"] + cur["to_km"]) / 2, 4)
    stretches.append(cur)
    for i, st in enumerate(stretches, 1):
        st["id"] = f"{id_prefix}{i:03d}"
    return stretches


def _side_totals(rows: list[dict]) -> dict:
    fill = sum(r["fill_m3"] for r in rows)
    cut = sum(r["cut_m3"] for r in rows)
    mass = sum(r["fill_mass_t"] for r in rows)
    cut_mass = sum(r["cut_mass_t"] for r in rows)
    return {
        "fill_m3": round(fill, 1),
        "cut_m3": round(cut, 1),
        "fill_mass_t": round(mass, 1),
        "cut_mass_t": round(cut_mass, 1),
    }


def _make_branch(bid: str, name: str, side: str, offset: int, rows, segs) -> dict:
    stretches = _merge_class_stretches(segs, id_prefix=f"{bid[0].upper()}S")
    totals = _side_totals(rows)
    return {
        "id": bid,
        "name": name,
        "side": side,
        "offset": offset,
        "station_count": len(rows),
        "segment_count": len(segs),
        "stretch_count": len(stretches),
        "fill_m3": totals["fill_m3"],
        "cut_m3": totals["cut_m3"],
        "net_m3": round(totals["fill_m3"] - totals["cut_m3"], 1),
        "fill_mass_t": totals["fill_mass_t"],
        "segments": segs,
        "stretches": stretches,
    }


def parse(path: str, cut_path: str | None = None) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    rhs = _read_side(wb["RHS"]) if "RHS" in wb.sheetnames else []
    lhs = _read_side(wb["LHS"]) if "LHS" in wb.sheetnames else []
    wb.close()

    if cut_path:
        cut_lut = _read_cut_lookup(cut_path)
        _apply_cut(rhs, cut_lut["rhs"])
        _apply_cut(lhs, cut_lut["lhs"])

    # RHS coordinates track chainage cleanly → anchor both sides to them.
    coords = _coord_lookup(rhs)
    if not coords:
        coords = _coord_lookup(lhs)

    rhs_segments = _aggregate(rhs, coords, id_prefix="R")
    lhs_segments = _aggregate(lhs, coords, id_prefix="L")

    rhs_branch = _make_branch("rhs", "RHS", "rhs", 1, rhs, rhs_segments)
    lhs_branch = _make_branch("lhs", "LHS", "lhs", -1, lhs, lhs_segments)
    branches = [lhs_branch, rhs_branch]

    # Combined per-100 m segment (both sides) for the profile/summary lists.
    combined: dict[int, dict] = {}
    for seg in [*lhs_segments, *rhs_segments]:
        idx = int(seg["from_km"] / AGG_KM + 1e-9)
        slot = combined.setdefault(
            idx,
            {
                "from_km": seg["from_km"],
                "to_km": seg["to_km"],
                "lat": seg["lat"],
                "lon": seg["lon"],
                "length_m": 0.0,
                "fill_m3": 0.0,
                "cut_m3": 0.0,
                "fill_mass_t": 0.0,
                "heights": [],
                "station_count": 0,
            },
        )
        slot["fill_m3"] += seg["fill_m3"]
        slot["cut_m3"] += seg["cut_m3"]
        slot["fill_mass_t"] += seg["fill_mass_t"]
        slot["length_m"] = max(slot["length_m"], seg["length_m"])
        slot["station_count"] += seg["station_count"]
        if seg["height_cl_m"] is not None:
            slot["heights"].append(seg["height_cl_m"])

    segments = []
    for i, idx in enumerate(sorted(combined), 1):
        c = combined[idx]
        h = sum(c["heights"]) / len(c["heights"]) if c["heights"] else None
        segments.append(
            {
                "id": f"CF{idx + 1:04d}",
                "from_km": c["from_km"],
                "to_km": c["to_km"],
                "mid_km": round((c["from_km"] + c["to_km"]) / 2, 4),
                "lat": c["lat"],
                "lon": c["lon"],
                "length_m": round(c["length_m"], 2),
                "fill_m3": round(c["fill_m3"], 1),
                "cut_m3": round(c["cut_m3"], 1),
                "net_m3": round(c["fill_m3"] - c["cut_m3"], 1),
                "fill_mass_t": round(c["fill_mass_t"], 1),
                "height_cl_m": round(h, 3) if h is not None else None,
                "station_count": c["station_count"],
                "class_id": _classify(h),
            }
        )
    stretches = _merge_class_stretches(segments, id_prefix="CFS")

    profile = [
        {
            "chainage_km": s["mid_km"],
            "fill_m3": s["fill_m3"],
            "cut_m3": s["cut_m3"],
            "net_m3": s["net_m3"],
            "height_cl_m": s["height_cl_m"],
            "class_id": s["class_id"],
            "tcs": None,
        }
        for s in segments
    ]

    lhs_t = _side_totals(lhs)
    rhs_t = _side_totals(rhs)
    total_fill = round(lhs_t["fill_m3"] + rhs_t["fill_m3"], 1)
    total_cut = round(lhs_t["cut_m3"] + rhs_t["cut_m3"], 1)
    total_mass = round(lhs_t["fill_mass_t"] + rhs_t["fill_mass_t"], 1)
    total_cut_mass = round(lhs_t["cut_mass_t"] + rhs_t["cut_mass_t"], 1)

    from_km = segments[0]["from_km"] if segments else None
    to_km = segments[-1]["to_km"] if segments else None

    summary = {
        "chainage_covered": (
            f"{from_km:.2f}–{to_km:.2f} km" if from_km is not None else None
        ),
        "right_of_way": "60 m (30 m each side)",
        "fill_m3": total_fill,
        "cut_m3": total_cut,
        "net_m3": round(total_fill - total_cut, 1),
        "fill_mt": total_mass,
        "cut_mt": total_cut_mass,
        "net_mt": round(total_mass - total_cut_mass, 1),
        "lhs_fill_m3": lhs_t["fill_m3"],
        "lhs_cut_m3": lhs_t["cut_m3"],
        "lhs_net_m3": round(lhs_t["fill_m3"] - lhs_t["cut_m3"], 1),
        "lhs_fill_mt": lhs_t["fill_mass_t"],
        "rhs_fill_m3": rhs_t["fill_m3"],
        "rhs_cut_m3": rhs_t["cut_m3"],
        "rhs_net_m3": round(rhs_t["fill_m3"] - rhs_t["cut_m3"], 1),
        "rhs_fill_mt": rhs_t["fill_mass_t"],
        "tcs_breakdown": [],
    }

    class_counts = Counter(s["class_id"] for s in segments)

    return {
        "title": "Cut & Fill Earthwork",
        "description": (
            "Embankment fill volume & mass along the corridor (LHS + RHS, "
            "60 m RoW), shown as per-side fill-height points every 100 m. "
            "Cut volumes merged from the earlier earthwork workbook."
        ),
        "unit_volume": "m3",
        "unit_mass": "MT",
        "density_t_per_m3": DENSITY_T_PER_M3,
        "from_km": from_km,
        "to_km": to_km,
        "station_count": len(lhs) + len(rhs),
        "segment_count": len(segments),
        "stretch_count": len(stretches),
        "rhs_station_count": len(rhs),
        "lhs_station_count": len(lhs),
        "summary": summary,
        "class_counts": dict(class_counts),
        "tcs_ranges": [],
        "branches": branches,
        "segments": segments,
        "stretches": stretches,
        "profile": profile,
    }


def main() -> None:
    src = None
    for cand in DOWNLOADS_CANDIDATES:
        if os.path.isfile(cand):
            os.makedirs(os.path.dirname(DEFAULT_SRC), exist_ok=True)
            shutil.copy2(cand, DEFAULT_SRC)
            src = DEFAULT_SRC
            break
    if src is None and os.path.isfile(DEFAULT_SRC):
        src = DEFAULT_SRC
    if src is None:
        raise SystemExit(f"Source not found. Expected {DEFAULT_SRC}")

    cut_src = None
    for cand in CUT_DOWNLOADS_CANDIDATES:
        if os.path.isfile(cand):
            shutil.copy2(cand, CUT_SRC)
            cut_src = CUT_SRC
            break
    if cut_src is None and os.path.isfile(CUT_SRC):
        cut_src = CUT_SRC

    payload = parse(src, cut_path=cut_src)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, separators=(",", ":"))

    size_kb = os.path.getsize(OUT) / 1024
    print(
        f"Wrote LHS {payload['lhs_station_count']} / RHS {payload['rhs_station_count']} "
        f"stations → branches {[b['id'] + ':' + str(b['segment_count']) for b in payload['branches']]} "
        f"· total fill {payload['summary']['fill_m3']:,} m³ "
        f"· total cut {payload['summary']['cut_m3']:,} m³ "
        f"({size_kb:.1f} KB) → {OUT}",
        flush=True,
    )


if __name__ == "__main__":
    main()
