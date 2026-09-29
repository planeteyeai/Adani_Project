"""Parse Design HFL workbook/CSV into frontend/public/design_hfl.json."""
from __future__ import annotations

import csv
import json
import os
import re

import openpyxl

ROOT = os.path.dirname(__file__)
DEFAULT_SRC = os.path.join(ROOT, "data", "HFL.xlsx")
OUT = os.path.join(ROOT, "..", "frontend", "public", "design_hfl.json")

HEADER_ALIASES = {
    "id": "id",
    "p": "id",
    "chainage_km": "chainage_km",
    "chainage": "chainage_km",
    "latitude": "latitude",
    "lat": "latitude",
    "longitude": "longitude",
    "lon": "longitude",
    "design_hfl_continuous_m": "design_hfl_continuous_m",
    "design_hfl": "design_hfl_continuous_m",
    "hfl": "design_hfl_continuous_m",
}


def _norm_header(value: object) -> str:
    key = re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")
    return HEADER_ALIASES.get(key, key)


def _row_from_mapping(mapping: dict[str, object]) -> dict | None:
    chainage = mapping.get("chainage_km")
    hfl = mapping.get("design_hfl_continuous_m")
    lat = mapping.get("latitude")
    lon = mapping.get("longitude")
    if chainage is None or hfl is None or lat is None or lon is None:
        return None
    try:
        chainage_f = float(chainage)
        hfl_f = float(hfl)
        lat_f = float(lat)
        lon_f = float(lon)
    except (TypeError, ValueError):
        return None
    row_id = mapping.get("id")
    parsed_id = None
    if row_id not in (None, ""):
        try:
            parsed_id = int(float(row_id))
        except (TypeError, ValueError):
            m = re.search(r"\d+", str(row_id))
            parsed_id = int(m.group()) if m else None
    return {
        "id": parsed_id,
        "chainage_km": round(chainage_f, 3),
        "latitude": round(lat_f, 8),
        "longitude": round(lon_f, 8),
        "design_hfl_continuous_m": round(hfl_f, 2),
    }


def parse_csv(path: str) -> list[dict]:
    rows: list[dict] = []
    with open(path, newline="", encoding="utf-8-sig") as f:
        sample = f.read(4096)
        f.seek(0)
        delimiter = "\t" if sample.count("\t") > sample.count(",") else ","
        reader = csv.DictReader(f, delimiter=delimiter)
        if not reader.fieldnames:
            return rows
        field_map = {_norm_header(name): name for name in reader.fieldnames}
        for raw in reader:
            mapped = {key: raw.get(src) for key, src in field_map.items() if src is not None}
            row = _row_from_mapping(mapped)
            if row:
                if row["id"] is None:
                    row["id"] = len(rows) + 1
                rows.append(row)
    return rows


def _looks_like_header(row: tuple) -> bool:
    if not row:
        return False
    text = " ".join(str(v or "").lower() for v in row)
    return "chainage" in text or "latitude" in text or "hfl" in text


def parse_xlsx(path: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    rows: list[dict] = []

    # Headerless layout used by HFL.xlsx: ID, Chainage_km, Lat, Lon, Design_HFL
    if not _looks_like_header(first):
        for values in ws.iter_rows(min_row=1, values_only=True):
            if not values or len(values) < 5:
                continue
            mapped = {
                "id": values[0],
                "chainage_km": values[1],
                "latitude": values[2],
                "longitude": values[3],
                "design_hfl_continuous_m": values[4],
            }
            row = _row_from_mapping(mapped)
            if row:
                if row["id"] is None:
                    row["id"] = len(rows) + 1
                rows.append(row)
        return rows

    field_map = {_norm_header(v): idx for idx, v in enumerate(first)}
    for values in ws.iter_rows(min_row=2, values_only=True):
        mapped = {
            key: values[idx] if idx is not None and idx < len(values) else None
            for key, idx in field_map.items()
        }
        row = _row_from_mapping(mapped)
        if row:
            if row["id"] is None:
                row["id"] = len(rows) + 1
            rows.append(row)
    return rows


def main() -> None:
    src = os.environ.get("DESIGN_HFL_SRC", DEFAULT_SRC)
    if not os.path.isfile(src):
        raise SystemExit(f"source not found: {src}")

    rows = parse_xlsx(src) if src.lower().endswith((".xlsx", ".xls")) else parse_csv(src)
    if not rows:
        raise SystemExit(f"no HFL rows parsed from {src}")

    rows.sort(key=lambda r: r["chainage_km"])
    for i, row in enumerate(rows, start=1):
        row["id"] = i

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(rows, f, separators=(",", ":"))

    hfl_vals = [r["design_hfl_continuous_m"] for r in rows]
    print(f"wrote {OUT}")
    print(f"parsed {len(rows)} rows from {src}")
    print(
        f"chainage {rows[0]['chainage_km']} – {rows[-1]['chainage_km']} km | "
        f"HFL {min(hfl_vals):.2f} – {max(hfl_vals):.2f} m"
    )


if __name__ == "__main__":
    main()
